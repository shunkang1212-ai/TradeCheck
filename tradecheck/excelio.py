"""Excel 读取底层。

只做一件事：把 .xlsx 读成「单元格 -> (公式原文 / 缓存值)」的矩阵，供上层按模板配置解释。

关键约束（任务书 R06、第 6 节）：
- 单元格值读两遍：data_only=False 拿公式原文，data_only=True 拿缓存结果。
  读到公式缓存不等于重新计算，必须标记 cached_available。
- 不执行宏、不解析外部链接、不写回工作簿。本模块对文件只读。
- 文件过大、损坏、加密或不是 xlsx 时给出可理解的中文提示。
"""

from __future__ import annotations

import hashlib
import os
import zipfile
from dataclasses import dataclass
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

#: 单文件最大 20 MB
MAX_FILE_BYTES = 20 * 1024 * 1024
#: 扫描范围上限，防止畸形文件把内存吃满
MAX_SCAN_ROWS = 5000
MAX_SCAN_COLS = 128
#: 明细行数上限
MAX_DETAIL_ROWS = 2000


class WorkbookReadError(Exception):
    """文件无法读取。message 面向最终用户，不含客户敏感正文。"""


@dataclass
class CellInfo:
    """一个单元格的读取结果。"""

    row: int
    column: int
    value: Any = None  # 非公式单元格的原文值；公式单元格为缓存值（可能为 None）
    formula: Optional[str] = None
    is_formula: bool = False
    has_cache: bool = False
    external: bool = False  # 公式里出现外部引用特征

    @property
    def empty(self) -> bool:
        return self.value is None or (isinstance(self.value, str) and self.value.strip() == "")


@dataclass
class SheetData:
    name: str
    max_row: int
    max_column: int
    cells: Dict[Tuple[int, int], CellInfo]

    def get(self, row: int, column: int) -> CellInfo:
        if row < 1 or column < 1:
            return CellInfo(row=row, column=column)
        return self.cells.get((row, column)) or CellInfo(row=row, column=column)

    def row_texts(self, row: int) -> List[str]:
        out: List[str] = []
        for c in range(1, self.max_column + 1):
            v = self.get(row, c).value
            if isinstance(v, str) and v.strip():
                out.append(v.strip())
        return out

    def find_rows_containing(self, texts: List[str]) -> List[int]:
        """返回同时包含全部给定文本（忽略大小写与空白）的行号列表。"""
        wanted = [_norm(t) for t in texts]
        hits: List[int] = []
        for r in range(1, self.max_row + 1):
            joined = _norm(" ".join(self.row_texts(r)))
            if all(w in joined for w in wanted):
                hits.append(r)
        return hits


@dataclass
class WorkbookData:
    path: str
    file_name: str
    sha256: str
    size_bytes: int
    sheets: List[SheetData]

    def sheet(self, name: Optional[str] = None) -> SheetData:
        if not self.sheets:
            raise WorkbookReadError("工作簿中没有任何工作表。")
        if name:
            target = _norm(name)
            for s in self.sheets:
                if _norm(s.name) == target:
                    return s
            raise WorkbookReadError(
                f"未找到工作表“{name}”。该文件中可用的工作表：{self.sheet_names_text()}。"
            )
        return self.sheets[0]

    def sheet_names_text(self) -> str:
        return "、".join(s.name for s in self.sheets) or "（无）"

    def has_sheet(self, name: str) -> bool:
        target = _norm(name)
        return any(_norm(s.name) == target for s in self.sheets)


def _norm(text: Any) -> str:
    return " ".join(str(text).split()).casefold()


def sha256_file(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _looks_like_external(formula: str) -> bool:
    low = formula.casefold()
    return any(token in low for token in ("[", "http://", "https://", "\\\\", ".xlsx!", ".xls!"))


def read_workbook(path: str, display_name: Optional[str] = None) -> WorkbookData:
    """读取一个 .xlsx，返回只读的数据结构。

    抛 WorkbookReadError 时消息可直接展示给用户。
    """
    if not os.path.isfile(path):
        raise WorkbookReadError(f"文件不存在：{os.path.basename(path)}")

    size = os.path.getsize(path)
    if size == 0:
        raise WorkbookReadError(f"文件为空（0 字节）：{os.path.basename(path)}")
    if size > MAX_FILE_BYTES:
        raise WorkbookReadError(
            f"文件过大：{os.path.basename(path)} 为 {size / 1024 / 1024:.1f} MB，"
            f"超过本工具上限 {MAX_FILE_BYTES // 1024 // 1024} MB。"
        )

    digest = sha256_file(path)

    # 在载入 XML 前限制解压体积；xlsx 的压缩后大小不足以约束内存。
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(i.file_size for i in entries) > 100 * 1024 * 1024:
                raise WorkbookReadError("工作簿解压规模超过上限，本次未读取，请拆分文件。")
    except zipfile.BadZipFile:
        raise WorkbookReadError("文件不是有效的 .xlsx 或已损坏，请另存为普通 .xlsx。")

    try:
        wb_formula = load_workbook(path, data_only=False, keep_links=False)
    except zipfile.BadZipFile:
        raise WorkbookReadError(
            f"无法读取 {os.path.basename(path)}：文件不是有效的 .xlsx（压缩包损坏）。"
            "请确认文件未损坏，并另存为 .xlsx 后重试。"
        )
    except InvalidFileException as exc:
        raise WorkbookReadError(
            f"无法读取 {os.path.basename(path)}：格式不受支持。"
            f"本工具只处理 .xlsx。若文件是 .xls / .csv / 受密码保护的工作簿，请先另存为普通 .xlsx。"
            f"（底层信息：{exc}）"
        )
    except Exception as exc:  # pragma: no cover - 兜底，避免堆栈吓到用户
        raise WorkbookReadError(
            f"无法读取 {os.path.basename(path)}：{type(exc).__name__}。"
            "若文件受密码保护或已损坏，请解除保护后重试。"
        )

    try:
        wb_cached = load_workbook(path, data_only=True, keep_links=False)
    except Exception:
        wb_cached = None

    sheets: List[SheetData] = []
    try:
        for idx, ws_formula in enumerate(wb_formula.worksheets):
            ws_cached = None
            if wb_cached is not None and idx < len(wb_cached.worksheets):
                ws_cached = wb_cached.worksheets[idx]

            max_row = int(ws_formula.max_row or 1)
            max_col = int(ws_formula.max_column or 1)
            if max_row > MAX_SCAN_ROWS or max_col > MAX_SCAN_COLS:
                raise WorkbookReadError(
                    f"工作表范围超过 {MAX_SCAN_ROWS} 行或 {MAX_SCAN_COLS} 列，本次停止读取，未截断数据后核对。"
                )

            cells: Dict[Tuple[int, int], CellInfo] = {}
            for r in range(1, max_row + 1):
                for c in range(1, max_col + 1):
                    src = ws_formula.cell(row=r, column=c)
                    raw = src.value
                    is_formula = isinstance(raw, str) and raw.startswith("=")

                    if is_formula:
                        cached = None
                        if ws_cached is not None:
                            cached = ws_cached.cell(row=r, column=c).value
                        cells[(r, c)] = CellInfo(
                            row=r,
                            column=c,
                            value=cached,
                            formula=raw,
                            is_formula=True,
                            has_cache=cached is not None,
                            external=_looks_like_external(raw),
                        )
                    else:
                        cells[(r, c)] = CellInfo(row=r, column=c, value=raw)

            sheets.append(
                SheetData(name=ws_formula.title, max_row=max_row, max_column=max_col, cells=cells)
            )
    finally:
        for wb in (wb_formula, wb_cached):
            if wb is not None:
                try:
                    wb.close()
                except Exception:
                    pass

    return WorkbookData(
        path=path,
        file_name=display_name or os.path.basename(path),
        sha256=digest,
        size_bytes=size,
        sheets=sheets,
    )
