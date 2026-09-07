"""按模板配置从工作簿里提取商业发票 / 装箱单的数据。

本模块只负责「把单元格搬出来」，不做任何一致性判断。
所有取数都带上单元格位置，供报告追溯（任务书 5.）。
"""

from __future__ import annotations

from typing import Any, Dict, List, Optional, Tuple

from .excelio import MAX_DETAIL_ROWS, CellInfo, SheetData, WorkbookData, WorkbookReadError
from .models import (
    FIELD_MISSING,
    FIELD_NEGATIVE,
    FIELD_NON_NUMERIC,
    FIELD_OK,
    FIELD_UNVERIFIED,
    STATUS_SKIP,
    STATUS_TODO,
    CellRef,
    DetailRow,
    DocumentExtraction,
    Field,
    Issue,
    RowSnapshot,
)
from .money import to_decimal

#: 被视为「额外费用结构」的标签关键词（中英对照）。命中即提示不支持该金额结构。
EXTRA_COST_KEYWORDS = [
    "discount", "rebate", "freight", "shipping", "insurance", "handling",
    "commission", "bank charge", "tax", "vat", "gst", "surcharge", "other charge",
    "折扣", "折扣额", "优惠", "运费", "运保费", "保险费", "手续费", "佣金",
    "银行费用", "银行手续费", "税", "税额", "增值税", "附加费", "其他费用",
]

#: 数值类字段
NUMBER_FIELDS = {"qty", "unit_price", "amount", "total"}
#: 文本类字段
TEXT_FIELDS = {"model", "spec", "unit", "doc_no", "order_ref", "invoice_ref", "currency"}


def _norm(text: Any) -> str:
    return " ".join(str(text).split()).casefold()


def _strip_colon(text: str) -> str:
    t = str(text).strip()
    while t.endswith((":", "：")):
        t = t[:-1].strip()
    return t


def _match_alias(header_text: str, aliases: List[str]) -> bool:
    """表头文本是否命中某个别名。先判整词相等，再判包含。"""
    h = _norm(_strip_colon(header_text))
    if not h:
        return False
    for alias in aliases:
        a = _norm(alias)
        if h == a:
            return True
    for alias in aliases:
        a = _norm(alias)
        if a and (a in h or h in a):
            return True
    return False


# ---------------------------------------------------------------- 字段构造


def _cell_ref(sheet: SheetData, file_name: str, row: int, col: int) -> CellRef:
    from openpyxl.utils import get_column_letter

    return CellRef(file_name=file_name, sheet=sheet.name, coord=f"{get_column_letter(col)}{row}")


def build_field(
    name: str,
    info: Optional[CellInfo],
    sheet: SheetData,
    file_name: str,
    row: int,
    col: int,
    kind: str = "text",
) -> Field:
    """把一个单元格包装成 Field，同时保留原文与规范化值。"""
    if info is None:
        return Field(name=name, state=FIELD_MISSING)

    raw = info.value
    empty = info.empty

    if kind == "number":
        number = to_decimal(raw)
    else:
        number = to_decimal(raw) if isinstance(raw, (int, float)) and not isinstance(raw, bool) else None

    if isinstance(raw, str):
        text = raw.strip()
    elif raw is None:
        text = ""
    elif isinstance(raw, bool):
        text = "TRUE" if raw else "FALSE"
    elif isinstance(raw, float) and raw == int(raw) and abs(raw) < 1e15:
        # 1000.0 应显示成 1000，避免报告中多出无意义的 ".0"
        text = str(int(raw))
    else:
        text = str(raw)

    cell = _cell_ref(sheet, file_name, row, col)

    field = Field(
        name=name,
        text=text,
        raw=raw,
        number=number,
        key=_norm(text),
        cell=cell,
        is_formula=info.is_formula,
        formula=info.formula,
        cached_available=info.has_cache,
    )

    if info.is_formula:
        # 公式缓存不是重新计算的结果，一律标记未验证（任务书 R06）
        field.state = FIELD_UNVERIFIED
        field.note = (
            f"公式：{info.formula}"
            + ("（缓存值可用但未验证）" if info.has_cache else "（无缓存值）")
            + ("；含外部引用特征，未解析" if info.external else "")
        )
        return field

    if empty:
        field.state = FIELD_MISSING
        return field

    if kind == "number":
        if number is None:
            field.state = FIELD_NON_NUMERIC
            field.note = f"不是数字：{text!r}" if text else "空值"
        elif number < 0:
            field.state = FIELD_NEGATIVE
            field.note = "负数，需人工确认业务含义"
        else:
            field.state = FIELD_OK
    else:
        field.state = FIELD_OK

    return field


def build_blank_field(name: str) -> Field:
    return Field(name=name, state=FIELD_MISSING)


def field_kind(name: str) -> str:
    return "number" if name in NUMBER_FIELDS else "text"


# ---------------------------------------------------------------- 工作表与字段定位


def _pick_sheet(
    wb: WorkbookData, doc_cfg: Dict[str, Any], file_name: str, issues: List[Issue]
) -> SheetData:
    names = (doc_cfg.get("sheet") or {}).get("names") or []
    for n in names:
        if wb.has_sheet(n):
            return wb.sheet(n)
    first = wb.sheets[0]
    if names:
        issues.append(
            Issue(
                code="sheet-fallback",
                message=(
                    f"未找到配置的工作表（{ '、'.join(names) }），"
                    f"已回退到第一个工作表“{first.name}”。请确认模板选择是否正确。"
                ),
                severity=STATUS_TODO,
            )
        )
    return first


def _find_label_row(sheet: SheetData, label: str) -> List[Tuple[int, int]]:
    """找出所有文本命中 label 的单元格（行, 列）。"""
    target = _norm(_strip_colon(label))
    hits: List[Tuple[int, int]] = []
    if not target:
        return hits
    for r in range(1, sheet.max_row + 1):
        for c in range(1, sheet.max_column + 1):
            v = sheet.get(r, c).value
            if not isinstance(v, str):
                continue
            t = _norm(_strip_colon(v))
            if not t:
                continue
            if t == target or (len(target) >= 2 and target in t):
                hits.append((r, c))
    return hits


def _extract_header_fields(
    sheet: SheetData, doc_cfg: Dict[str, Any], file_name: str, issues: List[Issue]
) -> Dict[str, Field]:
    out: Dict[str, Field] = {}
    for name, fcfg in (doc_cfg.get("fields") or {}).items():
        label = fcfg.get("label", "")
        offsets = fcfg.get("offsets") or [[0, 1]]
        found: Optional[Field] = None
        for (r, c) in _find_label_row(sheet, label):
            for dr, dc in offsets:
                rr, cc = r + int(dr), c + int(dc)
                info = sheet.get(rr, cc)
                if info.empty and not info.is_formula:
                    continue
                found = build_field(name, info, sheet, file_name, rr, cc, field_kind(name))
                break
            if found is not None:
                break
        if found is None:
            field = build_blank_field(name)
            if fcfg.get("required", False):
                issues.append(
                    Issue(
                        code=f"missing-header-{name}",
                        message=f"未找到必需字段“{label}”（工作表：{sheet.name}）。",
                        severity=STATUS_TODO,
                    )
                )
            out[name] = field
        else:
            out[name] = found
    return out


def _find_header_row(sheet: SheetData, table_cfg: Dict[str, Any]) -> Optional[int]:
    anchors = table_cfg.get("header_anchor") or []
    columns: Dict[str, List[str]] = table_cfg.get("columns") or {}

    candidates = sheet.find_rows_containing(anchors) if anchors else []
    if not candidates:
        # 退一步：只要能命中任意一个必需列标题也算
        required = table_cfg.get("required_columns") or []
        wanted = [columns[r][0] for r in required if columns.get(r)]
        candidates = sheet.find_rows_containing(wanted) if wanted else []

    best_row: Optional[int] = None
    best_score = 0
    for r in candidates:
        score = 0
        for logical, aliases in columns.items():
            for c in range(1, sheet.max_column + 1):
                v = sheet.get(r, c).value
                if isinstance(v, str) and _match_alias(v, aliases):
                    score += 1
                    break
        if score > best_score:
            best_score, best_row = score, r
    if best_row is not None and best_score >= 1:
        return best_row
    return None


def _map_columns(
    sheet: SheetData, header_row: int, table_cfg: Dict[str, Any]
) -> Dict[str, int]:
    columns: Dict[str, List[str]] = table_cfg.get("columns") or {}
    mapping: Dict[str, int] = {}
    used: set = set()
    # 先做整词精确匹配，保证同义词不抢列
    for logical, aliases in columns.items():
        for c in range(1, sheet.max_column + 1):
            if c in used:
                continue
            v = sheet.get(header_row, c).value
            if not isinstance(v, str):
                continue
            if _norm(_strip_colon(v)) in {_norm(a) for a in aliases}:
                mapping[logical] = c
                used.add(c)
                break
    for logical, aliases in columns.items():
        if logical in mapping:
            continue
        for c in range(1, sheet.max_column + 1):
            if c in used:
                continue
            v = sheet.get(header_row, c).value
            if isinstance(v, str) and _match_alias(v, aliases):
                mapping[logical] = c
                used.add(c)
                break
    return mapping


# ---------------------------------------------------------------- 明细与合计


def _is_total_row(sheet: SheetData, row: int, labels: List[str]) -> bool:
    for c in range(1, sheet.max_column + 1):
        v = sheet.get(row, c).value
        if not isinstance(v, str):
            continue
        t = _norm(_strip_colon(v))
        if not t:
            continue
        for label in labels:
            lb = _norm(label)
            if lb and (t == lb or t.startswith(lb)):
                return True
    return False


def _row_is_blank(sheet: SheetData, row: int, cols: List[int]) -> bool:
    for c in cols:
        if not sheet.get(row, c).empty or sheet.get(row, c).is_formula:
            return False
    return True


def extract_document(
    wb: WorkbookData, role: str, template: Dict[str, Any]
) -> DocumentExtraction:
    """从工作簿里提取一份商业发票或装箱单。"""
    doc_cfg = template["documents"][role]
    file_name = wb.file_name
    issues: List[Issue] = []
    sheet = _pick_sheet(wb, doc_cfg, file_name, issues)

    result = DocumentExtraction(
        role=role,
        role_label=doc_cfg.get("label", role),
        file_name=file_name,
        sha256=wb.sha256,
        template_id=str(template["template_id"]),
        sheet=sheet.name,
    )

    result.header = _extract_header_fields(sheet, doc_cfg, file_name, issues)

    table_cfg = doc_cfg.get("table") or {}
    columns: Dict[str, List[str]] = table_cfg.get("columns") or {}
    required_columns: List[str] = table_cfg.get("required_columns") or []
    header_row = _find_header_row(sheet, table_cfg)

    if header_row is None:
        issues.append(
            Issue(
                code="table-not-found",
                message=(
                    f"未能识别明细表头行（工作表：{sheet.name}）。"
                    f"配置要求的列标题：{ '、'.join(columns.get(c, [c])[0] for c in required_columns) }。"
                    "请确认所选模板与文件布局一致；本工具 v0.1 不支持任意布局自动识别。"
                ),
                severity=STATUS_SKIP,
            )
        )
        result.issues = issues
        return result

    colmap = _map_columns(sheet, header_row, table_cfg)
    result.available_columns = sorted(colmap.keys())
    missing_cols = [c for c in required_columns if c not in colmap]
    for c in missing_cols:
        issues.append(
            Issue(
                code=f"missing-column-{c}",
                message=(
                    f"明细表缺少必需列“{ '、'.join(columns.get(c, [c])) }”"
                    f"（工作表：{sheet.name}，表头第 {header_row} 行）。"
                ),
                severity=STATUS_SKIP,
            )
        )
    if required_columns and set(required_columns).issubset(set(colmap.keys())) is False:
        result.table_found = False
        result.issues = issues
        return result
    result.table_found = True

    stop_cfg = table_cfg.get("stop") or {}
    blank_limit = int(stop_cfg.get("blank_rows", 2) or 2)
    total_labels = [str(x) for x in (stop_cfg.get("total_labels") or [])]

    mapped_cols = [colmap[c] for c in colmap]
    rows: List[DetailRow] = []
    snapshots: List[RowSnapshot] = []
    blank_streak = 0
    truncated = False

    r = header_row + 1
    while r <= sheet.max_row:
        if total_labels and _is_total_row(sheet, r, total_labels):
            break
        if _row_is_blank(sheet, r, mapped_cols):
            blank_streak += 1
            if blank_streak >= blank_limit:
                break
            r += 1
            continue
        blank_streak = 0

        if len(rows) >= MAX_DETAIL_ROWS:
            truncated = True
            break

        # 额外费用/折扣/税费行不是产品明细，不参与数量匹配；
        # 它们由 _scan_extra_costs 记录，并在 R05 里提示「不支持该金额结构」。
        # 判定依据：该行第一个非空文本单元格是费用类关键词，且型号列为空。
        first_text = ""
        first_col = 0
        for c in range(1, sheet.max_column + 1):
            v = sheet.get(r, c).value
            if isinstance(v, str) and v.strip():
                first_text, first_col = v.strip(), c
                break
        model_cell = sheet.get(r, colmap.get("model", 1))
        if (
            first_text
            and any(k in _norm(_strip_colon(first_text)) for k in EXTRA_COST_KEYWORDS)
            and model_cell.empty
        ):
            issues.append(
                Issue(
                    code="extra-cost-row",
                    message=(
                        f"第 {r} 行“{first_text}”被识别为额外费用/折扣/税费行，"
                        "未作为产品明细参与数量核对。"
                    ),
                    cell=_cell_ref(sheet, file_name, r, first_col),
                    severity=STATUS_TODO,
                )
            )
            r += 1
            continue

        fields: Dict[str, Field] = {}
        for logical, col in colmap.items():
            # 空单元格也要走 build_field，否则会丢掉「该单元格其实是公式，只是没有缓存值」这一信息
            fields[logical] = build_field(
                logical, sheet.get(r, col), sheet, file_name, r, col, field_kind(logical)
            )

        rows.append(DetailRow(row_index=r, fields=fields))

        # 行快照：标签 + 同行数值，用于额外费用结构识别
        label_text = ""
        label_cell = None
        numbers: List[Dict[str, Any]] = []
        for c in range(1, sheet.max_column + 1):
            info = sheet.get(r, c)
            if info.empty:
                continue
            v = info.value
            if isinstance(v, str) and not label_text:
                label_text = v.strip()
                label_cell = _cell_ref(sheet, file_name, r, c)
            num = to_decimal(v)
            if num is not None and not isinstance(v, bool):
                numbers.append(
                    {
                        "text": str(v),
                        "number": num,
                        "cell": _cell_ref(sheet, file_name, r, c),
                    }
                )
        snapshots.append(
            RowSnapshot(row_index=r, label=label_text, label_cell=label_cell, numbers=numbers)
        )
        r += 1

    if truncated:
        raise WorkbookReadError(f"明细超过 {MAX_DETAIL_ROWS} 行，本次停止提取；请拆分后重新核对。")

    result.rows = rows
    result.row_snapshots = snapshots

    if not rows:
        issues.append(
            Issue(
                code="no-detail-rows",
                message=f"明细为空（工作表：{sheet.name}，表头第 {header_row} 行之后没有数据行）。",
                severity=STATUS_SKIP,
            )
        )

    result.total = _extract_total(
        sheet, doc_cfg.get("total"), header_row, colmap, file_name, issues
    )
    result.extra_cost_rows = _scan_extra_costs(sheet, header_row, file_name)
    result.issues = issues
    return result


def _scan_extra_costs(
    sheet: SheetData, header_row: int, file_name: str
) -> List[Dict[str, Any]]:
    """扫描表头行之后的所有行，找出形如「运费 / 折扣 / 税费」的标签 + 数值。

    v0.1 只识别并报告，不参与金额计算（任务书 R05：发现额外费用结构须提示，
    不能擅自忽略后宣称通过）。
    """
    hits: List[Dict[str, Any]] = []
    for r in range(header_row + 1, sheet.max_row + 1):
        label_text = ""
        label_col = 0
        for c in range(1, sheet.max_column + 1):
            v = sheet.get(r, c).value
            if isinstance(v, str) and v.strip() and not label_text:
                label_text = v.strip()
                label_col = c
                break
        if not label_text:
            continue
        low = _norm(_strip_colon(label_text))
        if not any(k in low for k in EXTRA_COST_KEYWORDS):
            continue
        numbers: List[Dict[str, Any]] = []
        for c in range(1, sheet.max_column + 1):
            if c == label_col:
                continue
            info = sheet.get(r, c)
            if info.empty or isinstance(info.value, bool):
                continue
            num = to_decimal(info.value)
            if num is None:
                continue
            numbers.append(
                {
                    "text": str(info.value),
                    "number": str(num),
                    "cell": _cell_ref(sheet, file_name, r, c).to_dict(),
                }
            )
        if numbers:
            hits.append(
                {
                    "label": label_text,
                    "row": r,
                    "label_cell": _cell_ref(sheet, file_name, r, label_col).to_dict(),
                    "numbers": numbers,
                }
            )
    return hits


def _extract_total(
    sheet: SheetData,
    total_cfg: Optional[Dict[str, Any]],
    header_row: int,
    colmap: Dict[str, int],
    file_name: str,
    issues: List[Issue],
) -> Optional[Field]:
    if not total_cfg:
        return None
    labels = [str(x) for x in (total_cfg.get("labels") or [])]
    value_col = colmap.get(total_cfg.get("value_from_column", "amount"))
    if not labels or value_col is None:
        return None

    candidates: List[int] = []
    for r in range(header_row + 1, sheet.max_row + 1):
        if _is_total_row(sheet, r, labels):
            candidates.append(r)

    for r in reversed(candidates):
        info = sheet.get(r, value_col)
        if not info.empty:
            return build_field("total", info, sheet, file_name, r, value_col, "number")
    if candidates:
        r = candidates[-1]
        return build_field("total", sheet.get(r, value_col), sheet, file_name, r, value_col, "number")
    return None


def extract_pair(
    invoice_wb: WorkbookData, packing_wb: WorkbookData, template: Dict[str, Any]
) -> Tuple[DocumentExtraction, DocumentExtraction]:
    inv = extract_document(invoice_wb, "invoice", template)
    pac = extract_document(packing_wb, "packing", template)
    return inv, pac
