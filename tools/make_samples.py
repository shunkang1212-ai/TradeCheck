"""生成虚构样例 .xlsx。

全部数据均为虚构：型号、订单号、金额都由本脚本编造，不含任何真实客户、
账号、银行信息或商业资料（任务书 6.）。

用法：
    python tools/make_samples.py
"""

from __future__ import annotations

import os
import shutil
import sys
import zipfile

from openpyxl import Workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAMPLES = os.path.join(ROOT, "samples")

sys.path.insert(0, ROOT)


# ---------------------------------------------------------------- 模板 A 写入器


def _title(ws, text: str) -> None:
    ws["A1"] = text


def write_invoice_a(path, doc_no, order_ref, currency, rows, extra_rows=None, total=None,
                    amount_formula=False, total_formula=False):
    """模板 A 商业发票。表头第 7 行，数据从第 8 行开始。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    _title(ws, "COMMERCIAL INVOICE")
    ws["A3"], ws["C3"] = "Invoice No.", doc_no
    ws["A4"], ws["C4"] = "Order Ref.", order_ref
    if currency is not None:
        ws["A5"], ws["C5"] = "Currency", currency

    headers = ["Item No.", "Model No.", "Specification", "Unit", "Quantity", "Unit Price", "Amount"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=7, column=i, value=h)

    r = 8
    for idx, row in enumerate(rows, start=1):
        model, spec, unit, qty, price = row[0], row[1], row[2], row[3], row[4]
        ws.cell(row=r, column=1, value=idx)
        ws.cell(row=r, column=2, value=model)
        ws.cell(row=r, column=3, value=spec)
        ws.cell(row=r, column=4, value=unit)
        ws.cell(row=r, column=5, value=qty)
        ws.cell(row=r, column=6, value=price)
        if amount_formula:
            ws.cell(row=r, column=7, value=f"=E{r}*F{r}")
        else:
            ws.cell(row=r, column=7, value=row[5] if len(row) > 5 else None)
        r += 1

    for label, value in (extra_rows or []):
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=7, value=value)
        r += 1

    ws.cell(row=r, column=1, value="Total")
    if total_formula:
        ws.cell(row=r, column=5, value=f"=SUM(E8:E{r-1})")
        ws.cell(row=r, column=7, value=f"=SUM(G8:G{r-1})")
    else:
        ws.cell(row=r, column=5, value=sum(x[3] for x in rows if isinstance(x[3], (int, float))))
        ws.cell(row=r, column=7, value=total)

    wb.save(path)
    return path


def write_packing_a(path, order_ref, invoice_ref, rows, total=None):
    """模板 A 装箱单。表头第 7 行，数据从第 8 行开始。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Packing List"
    _title(ws, "PACKING LIST")
    ws["A3"], ws["C3"] = "Order Ref.", order_ref
    ws["A4"], ws["C4"] = "Invoice Ref.", invoice_ref

    headers = ["Item No.", "Model No.", "Specification", "Unit", "Quantity"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=7, column=i, value=h)

    r = 8
    for idx, row in enumerate(rows, start=1):
        model, spec, unit, qty = row[0], row[1], row[2], row[3]
        ws.cell(row=r, column=1, value=idx)
        ws.cell(row=r, column=2, value=model)
        ws.cell(row=r, column=3, value=spec)
        ws.cell(row=r, column=4, value=unit)
        ws.cell(row=r, column=5, value=qty)
        r += 1

    ws.cell(row=r, column=1, value="Total")
    if total is None:
        total = sum(x[3] for x in rows if isinstance(x[3], (int, float)))
    ws.cell(row=r, column=5, value=total)
    wb.save(path)
    return path


# ---------------------------------------------------------------- 模板 B 写入器


def write_invoice_b(path, doc_no, order_ref, currency, rows, extra_rows=None, total=None):
    """模板 B 商业发票：抬头值在标签右侧一列，表头第 7 行。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "商业发票"
    _title(ws, "商业发票（虚构样例）")
    ws["A3"], ws["B3"] = "发票号码", doc_no
    ws["A4"], ws["B4"] = "订单编号", order_ref
    if currency is not None:
        ws["A5"], ws["B5"] = "币种", currency

    headers = ["序号", "型号", "规格", "单位", "数量", "单价", "金额"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=7, column=i, value=h)

    r = 8
    for idx, row in enumerate(rows, start=1):
        ws.cell(row=r, column=1, value=idx)
        ws.cell(row=r, column=2, value=row[0])
        ws.cell(row=r, column=3, value=row[1])
        ws.cell(row=r, column=4, value=row[2])
        ws.cell(row=r, column=5, value=row[3])
        ws.cell(row=r, column=6, value=row[4])
        ws.cell(row=r, column=7, value=row[5])
        r += 1

    for label, value in (extra_rows or []):
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=7, value=value)
        r += 1

    ws.cell(row=r, column=1, value="合计")
    ws.cell(row=r, column=5, value=sum(x[3] for x in rows if isinstance(x[3], (int, float))))
    ws.cell(row=r, column=7, value=total)
    wb.save(path)
    return path


def write_packing_b(path, order_ref, invoice_ref, rows, total=None):
    """模板 B 装箱单：列顺序不同（序号/箱号/型号/规格/单位/数量），表头第 6 行。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "装箱单"
    _title(ws, "装箱单（虚构样例）")
    ws["A3"], ws["B3"] = "订单编号", order_ref
    ws["A4"], ws["B4"] = "发票号码", invoice_ref

    headers = ["序号", "箱号", "型号", "规格", "单位", "数量"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=6, column=i, value=h)

    r = 7
    for idx, row in enumerate(rows, start=1):
        model, spec, unit, qty = row[0], row[1], row[2], row[3]
        carton = row[4] if len(row) > 4 else f"C{idx:03d}"
        ws.cell(row=r, column=1, value=idx)
        ws.cell(row=r, column=2, value=carton)
        ws.cell(row=r, column=3, value=model)
        ws.cell(row=r, column=4, value=spec)
        ws.cell(row=r, column=5, value=unit)
        ws.cell(row=r, column=6, value=qty)
        r += 1

    ws.cell(row=r, column=1, value="合计")
    if total is None:
        total = sum(x[3] for x in rows if isinstance(x[3], (int, float)))
    ws.cell(row=r, column=6, value=total)
    wb.save(path)
    return path


# ---------------------------------------------------------------- 公式缓存注入


def inject_stale_cache(path: str, replacements: dict) -> None:
    """给公式单元格塞入一个「过期缓存值」，用于验证缓存不等于重新计算。

    replacements: {"G8": "9999", ...}  -> <f>...</f> 后面插入 <v>9999</v>
    """
    tmp = path + ".tmp"
    with zipfile.ZipFile(path, "r") as zin:
        names = zin.namelist()
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for name in names:
                data = zin.read(name)
                if name == "xl/worksheets/sheet1.xml":
                    xml = data.decode("utf-8")
                    for coord, value in replacements.items():
                        token = f'<c r="{coord}"'
                        start = xml.find(token)
                        if start < 0:
                            continue
                        end = xml.find("</c>", start)
                        if end < 0:
                            continue
                        seg = xml[start:end]
                        if "</f>" in seg:
                            seg2 = seg.replace("</f>", f"</f><v>{value}</v>", 1)
                            xml = xml[:start] + seg2 + xml[end:]
                    data = xml.encode("utf-8")
                zout.writestr(name, data)
    shutil.move(tmp, path)


def write_unsupported(path):
    """一个两种模板都不匹配的布局，用于验证「不支持的模板」提示。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "随便一个表"
    ws["A3"] = "客户名称"
    ws["B3"] = "某虚构公司"
    ws["A5"] = "备注"
    ws["B5"] = "这是一个不受支持的布局，用于验证提示"
    wb.save(path)
    return path


# ---------------------------------------------------------------- 样例定义


def build_all() -> None:
    a_dir = os.path.join(SAMPLES, "A")
    b_dir = os.path.join(SAMPLES, "B")
    os.makedirs(a_dir, exist_ok=True)
    os.makedirs(b_dir, exist_ok=True)

    # 1. 正常单据
    write_invoice_a(
        os.path.join(a_dir, "normal_invoice.xlsx"),
        "TCI-2026-0001", "PO-88421", "USD",
        [
            ("TC-100", "RED-L", "PCS", 500, 2.50, 1250.00),
            ("TC-200", "BLU-M", "PCS", 300, 4.00, 1200.00),
            ("TC-300", "", "SET", 20, 35.00, 700.00),
        ],
        total=3150.00,
    )
    write_packing_a(
        os.path.join(a_dir, "normal_packing.xlsx"),
        "PO-88421", "TCI-2026-0001",
        [
            ("TC-100", "RED-L", "PCS", 500),
            ("TC-200", "BLU-M", "PCS", 300),
            ("TC-300", "", "SET", 20),
        ],
    )

    # 2. 发票 1000 / 装箱单 980
    write_invoice_a(
        os.path.join(a_dir, "qty_short_invoice.xlsx"),
        "TCI-2026-0002", "PO-88422", "USD",
        [("TC-100", "RED-L", "PCS", 1000, 2.50, 2500.00)],
        total=2500.00,
    )
    write_packing_a(
        os.path.join(a_dir, "qty_short_packing.xlsx"),
        "PO-88422", "TCI-2026-0002",
        [("TC-100", "RED-L", "PCS", 980)],
    )

    # 3. 装箱单同一型号拆成多行
    write_invoice_a(
        os.path.join(a_dir, "split_invoice.xlsx"),
        "TCI-2026-0003", "PO-88423", "USD",
        [("TC-100", "RED-L", "PCS", 1000, 1.20, 1200.00)],
        total=1200.00,
    )
    write_packing_a(
        os.path.join(a_dir, "split_packing.xlsx"),
        "PO-88423", "TCI-2026-0003",
        [
            ("TC-100", "RED-L", "PCS", 600),
            ("TC-100", "RED-L", "PCS", 400),
        ],
    )

    # 4. 同型号不同规格（若错误合并会漏报）
    write_invoice_a(
        os.path.join(a_dir, "spec_invoice.xlsx"),
        "TCI-2026-0004", "PO-88424", "USD",
        [
            ("TC-100", "RED-L", "PCS", 500, 2.00, 1000.00),
            ("TC-100", "BLU-M", "PCS", 300, 2.00, 600.00),
        ],
        total=1600.00,
    )
    write_packing_a(
        os.path.join(a_dir, "spec_packing.xlsx"),
        "PO-88424", "TCI-2026-0004",
        [
            ("TC-100", "RED-L", "PCS", 300),
            ("TC-100", "BLU-M", "PCS", 500),
        ],
    )

    # 5. 明细顺序不同
    write_invoice_a(
        os.path.join(a_dir, "reorder_invoice.xlsx"),
        "TCI-2026-0005", "PO-88425", "USD",
        [
            ("AA-01", "S1", "PCS", 10, 1.00, 10.00),
            ("BB-02", "S2", "PCS", 20, 2.00, 40.00),
            ("CC-03", "S3", "PCS", 30, 3.00, 90.00),
        ],
        total=140.00,
    )
    write_packing_a(
        os.path.join(a_dir, "reorder_packing.xlsx"),
        "PO-88425", "TCI-2026-0005",
        [
            ("CC-03", "S3", "PCS", 30),
            ("AA-01", "S1", "PCS", 10),
            ("BB-02", "S2", "PCS", 20),
        ],
    )

    # 6. 件 vs 箱（单位不同，不比较数量）
    write_invoice_a(
        os.path.join(a_dir, "unit_invoice.xlsx"),
        "TCI-2026-0006", "PO-88426", "USD",
        [("TC-700", "STD", "PCS", 1000, 1.00, 1000.00)],
        total=1000.00,
    )
    write_packing_a(
        os.path.join(a_dir, "unit_packing.xlsx"),
        "PO-88426", "TCI-2026-0006",
        [("TC-700", "STD", "CTN", 1000)],
    )

    # 7. 行金额错误（第一行错，第二行对）
    write_invoice_a(
        os.path.join(a_dir, "amount_invoice.xlsx"),
        "TCI-2026-0007", "PO-88427", "USD",
        [
            ("TC-500", "STD", "PCS", 120, 12.50, 1250.00),
            ("TC-600", "STD", "PCS", 40, 3.00, 120.00),
        ],
        total=1370.00,
    )
    write_packing_a(
        os.path.join(a_dir, "amount_packing.xlsx"),
        "PO-88427", "TCI-2026-0007",
        [
            ("TC-500", "STD", "PCS", 120),
            ("TC-600", "STD", "PCS", 40),
        ],
    )

    # 8. 小数舍入边界（HALF_UP + Decimal 精确性）
    write_invoice_a(
        os.path.join(a_dir, "round_invoice.xlsx"),
        "TCI-2026-0008", "PO-88428", "USD",
        [
            ("TC-R1", "", "PCS", 1, 0.125, 0.13),
            ("TC-R2", "", "PCS", 1, 2.675, 2.68),
            ("TC-R3", "", "PCS", 3, 0.10, 0.30),
        ],
        total=3.11,
    )
    write_packing_a(
        os.path.join(a_dir, "round_packing.xlsx"),
        "PO-88428", "TCI-2026-0008",
        [
            ("TC-R1", "", "PCS", 1),
            ("TC-R2", "", "PCS", 1),
            ("TC-R3", "", "PCS", 3),
        ],
    )

    # 9. 缺少字段（无币种 / 型号缺失 / 数量非数字）
    write_invoice_a(
        os.path.join(a_dir, "missing_invoice.xlsx"),
        "TCI-2026-0009", "PO-88429", None,
        [
            ("TC-700", "STD", "PCS", 10, 5.00, 50.00),
            ("", "STD", "PCS", 10, 5.00, 50.00),
            ("TC-800", "STD", "PCS", "N/A", 5.00, 50.00),
        ],
        total=150.00,
    )
    write_packing_a(
        os.path.join(a_dir, "missing_packing.xlsx"),
        "PO-88429", "TCI-2026-0009",
        [
            ("TC-700", "STD", "PCS", 10),
            ("TC-800", "STD", "PCS", 10),
        ],
    )

    # 10a. 公式无缓存值
    p = write_invoice_a(
        os.path.join(a_dir, "formula_nocache_invoice.xlsx"),
        "TCI-2026-0010", "PO-88430", "USD",
        [
            ("TC-900", "STD", "PCS", 100, 2.00, None),
            ("TC-901", "STD", "PCS", 50, 4.00, None),
        ],
        total=None, amount_formula=True, total_formula=True,
    )
    write_packing_a(
        os.path.join(a_dir, "formula_nocache_packing.xlsx"),
        "PO-88430", "TCI-2026-0010",
        [
            ("TC-900", "STD", "PCS", 100),
            ("TC-901", "STD", "PCS", 50),
        ],
    )

    # 10b. 公式有缓存值，但缓存是过期的
    p = write_invoice_a(
        os.path.join(a_dir, "formula_stale_invoice.xlsx"),
        "TCI-2026-0011", "PO-88431", "USD",
        [
            ("TC-900", "STD", "PCS", 100, 2.00, None),
            ("TC-901", "STD", "PCS", 50, 4.00, None),
        ],
        total=None, amount_formula=True, total_formula=True,
    )
    inject_stale_cache(p, {"G8": "9999", "G9": "8888", "G10": "7777"})
    write_packing_a(
        os.path.join(a_dir, "formula_stale_packing.xlsx"),
        "PO-88431", "TCI-2026-0011",
        [
            ("TC-900", "STD", "PCS", 100),
            ("TC-901", "STD", "PCS", 50),
        ],
    )

    # 11. 额外费用结构（Freight）+ 总额不符
    write_invoice_a(
        os.path.join(a_dir, "extracost_invoice.xlsx"),
        "TCI-2026-0012", "PO-88432", "USD",
        [("TC-1000", "STD", "PCS", 100, 10.00, 1000.00)],
        extra_rows=[("Freight", 120.00)],
        total=1120.00,
    )
    write_packing_a(
        os.path.join(a_dir, "extracost_packing.xlsx"),
        "PO-88432", "TCI-2026-0012",
        [("TC-1000", "STD", "PCS", 100)],
    )

    # 11b. 未知币种
    write_invoice_a(
        os.path.join(a_dir, "unknownccy_invoice.xlsx"),
        "TCI-2026-0013", "PO-88433", "XXX",
        [("TC-1100", "STD", "PCS", 10, 5.00, 50.00)],
        total=50.00,
    )
    write_packing_a(
        os.path.join(a_dir, "unknownccy_packing.xlsx"),
        "PO-88433", "TCI-2026-0013",
        [("TC-1100", "STD", "PCS", 10)],
    )

    # 12. 型号前导零：两边都是 "007-A" 时必须原样保留并匹配上
    write_invoice_a(
        os.path.join(a_dir, "zero_ok_invoice.xlsx"),
        "TCI-2026-0014", "PO-88434", "USD",
        [("007-A", "STD", "PCS", 100, 1.00, 100.00)],
        total=100.00,
    )
    write_packing_a(
        os.path.join(a_dir, "zero_ok_packing.xlsx"),
        "PO-88434", "TCI-2026-0014",
        [("007-A", "STD", "PCS", 100)],
    )

    # 13. 型号 "007-A" 与 "7-A" 不是同一个产品，不能因数值化而混为一谈
    write_invoice_a(
        os.path.join(a_dir, "zero_mismatch_invoice.xlsx"),
        "TCI-2026-0015", "PO-88435", "USD",
        [("007-A", "STD", "PCS", 100, 1.00, 100.00)],
        total=100.00,
    )
    write_packing_a(
        os.path.join(a_dir, "zero_mismatch_packing.xlsx"),
        "PO-88435", "TCI-2026-0015",
        [("7-A", "STD", "PCS", 100)],
    )

    write_unsupported(os.path.join(SAMPLES, "unsupported_layout.xlsx"))

    # ---- 模板 B（中文表头）----
    write_invoice_b(
        os.path.join(b_dir, "normal_invoice.xlsx"),
        "TCB-2026-0001", "PO-90001", "CNY",
        [
            ("M-1001", "黑色", "件", 200, 15.00, 3000.00),
            ("M-1002", "白色", "件", 150, 22.00, 3300.00),
            ("M-2005", "", "套", 30, 120.00, 3600.00),
        ],
        total=9900.00,
    )
    write_packing_b(
        os.path.join(b_dir, "normal_packing.xlsx"),
        "PO-90001", "TCB-2026-0001",
        [
            ("M-1001", "黑色", "件", 200),
            ("M-1002", "白色", "件", 150),
            ("M-2005", "", "套", 30),
        ],
    )

    write_invoice_b(
        os.path.join(b_dir, "split_invoice.xlsx"),
        "TCB-2026-0002", "PO-90002", "CNY",
        [("M-3001", "蓝色", "件", 900, 8.00, 7200.00)],
        total=7200.00,
    )
    write_packing_b(
        os.path.join(b_dir, "split_packing.xlsx"),
        "PO-90002", "TCB-2026-0002",
        [
            ("M-3001", "蓝色", "件", 500),
            ("M-3001", "蓝色", "件", 300),
            ("M-3001", "蓝色", "件", 100),
        ],
    )

    # 模板 B 短装场景（中文环境也验证一次数量差异）
    write_invoice_b(
        os.path.join(b_dir, "qty_short_invoice.xlsx"),
        "TCB-2026-0003", "PO-90003", "CNY",
        [("M-4001", "灰色", "件", 1000, 3.00, 3000.00)],
        total=3000.00,
    )
    write_packing_b(
        os.path.join(b_dir, "qty_short_packing.xlsx"),
        "PO-90003", "TCB-2026-0003",
        [("M-4001", "灰色", "件", 950)],
    )


if __name__ == "__main__":
    build_all()
    count = sum(len(f) for _, _, f in os.walk(SAMPLES))
    print(f"样例已生成：{SAMPLES}（{count} 个文件）")
