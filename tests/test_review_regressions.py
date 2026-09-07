"""复核中发现的业务与会话缺陷；预期值独立写明。"""
import base64
import copy
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from tradecheck.corrections import set_field_text
from tradecheck.excelio import WorkbookReadError, read_workbook
from tradecheck.models import Correction, FIELD_UNVERIFIED, STATUS_DIFF, STATUS_OK, STATUS_TODO
from tradecheck.money import parse_currency, to_decimal
from tradecheck.pipeline import PipelineError, check_pair, extract_pair
from tradecheck.web import server

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "samples" / "A"


def pair(name="qty_short"):
    return extract_pair(str(SAMPLES / f"{name}_invoice.xlsx"), str(SAMPLES / f"{name}_packing.xlsx"), "A")


class CurrencyRegression(unittest.TestCase):
    def test_complete_chinese_currency_names(self):
        for text, expected in [("欧元", "EUR"), ("日元", "JPY"), ("港元", "HKD"),
                               ("新加坡元", "SGD"), ("澳元", "AUD"), ("美元 USD", "USD")]:
            with self.subTest(text=text):
                self.assertEqual(parse_currency(text)[0], expected)

    def test_conflicting_or_ambiguous_currency_not_accepted(self):
        for text in ("USD/CNY", "美元 CNY", "¥", "￥", "$", "元", "欧元 USD"):
            with self.subTest(text=text):
                self.assertIsNone(parse_currency(text)[0])

    def test_decimal_comma_not_silently_multiplied(self):
        for text in ("1,25", "12,34", "1.234,50"):
            self.assertIsNone(to_decimal(text))


class RuleRegression(unittest.TestCase):
    def test_empty_correction_is_applied_and_audited(self):
        inv, pac, tpl = pair()
        res, _ = check_pair(inv, pac, tpl, [Correction("invoice:row.8.qty", "1000", "")])
        self.assertEqual(res.corrections[0].new_text, "")
        self.assertFalse(any(f.rule_id in ("R03", "R04") and f.status == STATUS_OK for f in res.findings))
        self.assertEqual(inv.rows[0].get("qty").text, "1000")

    def test_total_correction_uses_numeric_parsing(self):
        inv, pac, tpl = pair()
        res, _ = check_pair(inv, pac, tpl, [Correction("invoice:total", "2500", "2400")])
        f = next(f for f in res.findings if f.rule_id == "R05")
        self.assertEqual(f.status, STATUS_DIFF)
        self.assertIn("2,400.00", f.actual)

    def test_invalid_or_negative_total_is_pending_without_crash(self):
        for value in ("待定", "-1"):
            with self.subTest(value=value):
                inv, pac, tpl = pair()
                set_field_text(inv.total, value)
                res, _ = check_pair(inv, pac, tpl)
                self.assertTrue(any(f.rule_id == "R05" and f.status == STATUS_TODO for f in res.findings))

    def test_same_mixed_unit_sets_are_not_summed(self):
        inv, pac, tpl = pair()
        for doc in (inv, pac):
            second = copy.deepcopy(doc.rows[0])
            second.row_index += 1
            set_field_text(second.get("unit"), "BOX")
            doc.rows.append(second)
        res, _ = check_pair(inv, pac, tpl)
        self.assertFalse(any(f.rule_id == "R03" and f.status == STATUS_OK for f in res.findings))
        self.assertTrue(any("单位" in f.item and f.status == STATUS_TODO for f in res.findings))

    def test_missing_unit_in_one_split_row_blocks_comparison(self):
        inv, pac, tpl = pair("split")
        set_field_text(pac.rows[1].get("unit"), "")
        res, _ = check_pair(inv, pac, tpl)
        self.assertFalse(any(f.rule_id == "R03" and f.status == STATUS_OK for f in res.findings))

    def test_one_sided_spec_is_not_merged_into_quantity_pass(self):
        inv, pac, tpl = pair("spec")
        for row in pac.rows:
            set_field_text(row.get("spec"), "")
        res, _ = check_pair(inv, pac, tpl)
        self.assertTrue(any(f.rule_id == "R02" and f.status == STATUS_TODO for f in res.findings))
        self.assertFalse(any(f.rule_id == "R03" and f.status == STATUS_OK for f in res.findings))

    def test_formula_spec_blocks_matching(self):
        inv, pac, tpl = pair()
        inv.rows[0].get("spec").state = FIELD_UNVERIFIED
        res, _ = check_pair(inv, pac, tpl)
        self.assertFalse(any(f.rule_id == "R03" and f.status == STATUS_OK for f in res.findings))

    def test_extra_cost_not_passed_even_if_total_happens_to_equal(self):
        inv, pac, tpl = pair("extracost")
        set_field_text(inv.total, str(sum(r.get("amount").number for r in inv.rows)))
        res, _ = check_pair(inv, pac, tpl)
        self.assertFalse(any(f.rule_id == "R05" and f.status == STATUS_OK for f in res.findings))

    def test_zero_tolerance_is_respected(self):
        inv, pac, tpl = pair()
        tpl["money"]["tolerance"]["absolute"] = "0"
        set_field_text(inv.rows[0].get("amount"), "2500.005")
        res, _ = check_pair(inv, pac, tpl)
        self.assertTrue(any(f.rule_id == "R04" and f.status == STATUS_DIFF for f in res.findings))

    def test_unit_price_precision_preserved_in_evidence(self):
        inv, pac, tpl = pair("round")
        res, _ = check_pair(inv, pac, tpl)
        self.assertTrue(any("0.125" in f.evidence for f in res.findings if f.rule_id == "R04"))

    def test_corrected_missing_header_removes_obsolete_issue(self):
        inv, pac, tpl = pair("missing")
        res, _ = check_pair(inv, pac, tpl, [Correction("invoice:header.currency", "", "USD")])
        self.assertFalse(any("missing-header-currency" in f.item for f in res.findings))


class WorkbookBoundsRegression(unittest.TestCase):
    def test_excess_rows_rejected_without_silent_truncation(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "oversize.xlsx"
            wb = load_workbook(SAMPLES / "normal_invoice.xlsx")
            wb.active.cell(5001, 1, "不能静默遗漏")
            wb.save(p)
            wb.close()
            with self.assertRaisesRegex(WorkbookReadError, "停止读取"):
                read_workbook(str(p))

    def test_display_name_is_used_for_all_sources(self):
        inv, pac, tpl = extract_pair(str(SAMPLES / "qty_short_invoice.xlsx"), str(SAMPLES / "qty_short_packing.xlsx"), "A", "发票原文件.xlsx", "装箱原文件.xlsx")
        res, _ = check_pair(inv, pac, tpl)
        self.assertEqual(inv.file_name, "发票原文件.xlsx")
        self.assertEqual({s.file_name for f in res.findings for s in f.sources}, {"发票原文件.xlsx", "装箱原文件.xlsx"})


class SessionRegression(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.patches = [patch.object(server, "OUTPUT_DIR", self.tmp.name),
                        patch.object(server, "SESSIONS_DIR", str(Path(self.tmp.name) / "sessions"))]
        for p in self.patches:
            p.start()
        server._clear_session()
        self.payload = {
            "template": "A", "invoice_name": "实际发票.xlsx", "packing_name": "实际装箱单.xlsx",
            "invoice_b64": base64.b64encode((SAMPLES / "qty_short_invoice.xlsx").read_bytes()).decode(),
            "packing_b64": base64.b64encode((SAMPLES / "qty_short_packing.xlsx").read_bytes()).decode(),
        }

    def tearDown(self):
        server._clear_session()
        for p in reversed(self.patches):
            p.stop()
        self.tmp.cleanup()

    def test_upload_preserves_original_filename(self):
        data = server.do_extract(self.payload)
        self.assertEqual(data["invoice"]["file_name"], "实际发票.xlsx")
        self.assertEqual(data["packing"]["rows"][0]["fields"]["qty"]["file"], "实际装箱单.xlsx")

    def test_previous_tab_cannot_check_replaced_session(self):
        first = server.do_extract(self.payload)
        server.do_extract(self.payload)
        with self.assertRaises(PipelineError):
            server.do_check({"session_id": first["session_id"]})

    def test_failed_extraction_clears_previous_session_and_temp(self):
        first = server.do_extract(self.payload)
        bad = dict(self.payload, packing_b64=base64.b64encode(b"broken workbook").decode())
        with self.assertRaises(PipelineError):
            server.do_extract(bad)
        with self.assertRaises(PipelineError):
            server.do_check({"session_id": first["session_id"]})
        self.assertEqual(list(Path(server.SESSIONS_DIR).iterdir()), [])

    def test_same_second_reports_do_not_overwrite(self):
        data = server.do_extract(self.payload)
        payload = {"session_id": data["session_id"]}
        first = server.do_report(payload)
        second = server.do_report(payload)
        self.assertNotEqual(first["path"], second["path"])
        self.assertEqual(Path(first["path"]).read_text(), first["html"])

    def test_report_records_correction_and_original_source(self):
        data = server.do_extract(self.payload)
        report = server.do_report({"session_id": data["session_id"], "corrections": [{"target": "invoice:row.8.qty", "new_text": "980"}]})
        self.assertIn("实际发票.xlsx", report["html"])
        self.assertEqual(report["result"]["corrections"][0]["new_text"], "980")
        self.assertIn("基于用户修正后的数据", report["html"])


if __name__ == "__main__":
    unittest.main()
